"""Read the three local Excel workbooks into canonical production models."""

from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .models import Agreement, EvidenceBundle, InvestigationRequest, Payment, Transaction
from .security import canonical_checksum


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(_text(value)[:10])


def _decimal(value: Any) -> Decimal:
    return Decimal(str(value or 0))


def _optional_bool(value: Any) -> bool | None:
    normalized = _text(value).casefold()
    if not normalized:
        return None
    if normalized in {"yes", "true", "1"}:
        return True
    if normalized in {"no", "false", "0"}:
        return False
    raise ValueError(f"invalid boolean value: {value}")


class SpreadsheetEvidenceProvider:
    """Laptop-only adapter. It performs exact tenant/customer/quarter filtering."""

    def __init__(self, data_dir: Path) -> None:
        self.data_dir = data_dir

    def _rows(self, filename: str) -> list[dict[str, Any]]:
        workbook = load_workbook(self.data_dir / filename, read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = [_text(value) for value in next(iterator)]
        rows = [dict(zip(headers, row)) for row in iterator]
        workbook.close()
        return rows

    def tenant_for_customer(self, customer: str) -> str:
        matches = {
            _text(row["Tenant ID"])
            for row in self._rows("rebate_agreements.xlsx")
            if _text(row["Customer"]).casefold() == customer.strip().casefold()
        }
        if len(matches) != 1:
            raise LookupError("customer was not found in exactly one local tenant")
        return matches.pop()

    def retrieve(self, request: InvestigationRequest) -> EvidenceBundle:
        def belongs(row: dict[str, Any], include_quarter: bool = False) -> bool:
            base = (
                _text(row["Tenant ID"]) == request.tenant_id
                and _text(row["Customer"]).casefold() == request.customer.casefold()
            )
            return base and (not include_quarter or _text(row["Quarter"]) == request.quarter)

        agreement_rows = [row for row in self._rows("rebate_agreements.xlsx") if belongs(row)]
        transaction_rows = [row for row in self._rows("transactions.xlsx") if belongs(row, True)]
        payment_rows = [row for row in self._rows("rebate_payments.xlsx") if belongs(row, True)]

        agreements = [Agreement(
            id=_text(row["Agreement ID"]), version=_text(row["Agreement Version"]),
            effective_from=_date(row["Effective Date"]), effective_to=_date(row["Expiration Date"]),
            rate=_decimal(row["Rebate Rate"]), threshold=_decimal(row["Quarterly Minimum Sales"]),
            eligible_products={item for item in _text(row["Eligible Product Codes"]).split(";") if item},
            status=_text(row["Status"]).casefold(), currency=_text(row["Currency"]),
            ambiguous=_text(row["Ambiguity Flag"]).casefold() == "yes",
            amendment_priority=int(row["Amendment Priority"] or 0),
            supersedes_agreement_id=_text(row["Supersedes Agreement ID"]) or None,
            human_review_threshold=_decimal(row["Human Review Threshold"]) if row["Human Review Threshold"] is not None else None,
            source_checksum=canonical_checksum(row),
        ) for row in agreement_rows]

        transactions = [Transaction(
            id=_text(row["Transaction ID"]), source_version=_text(row["Source Version"]),
            transaction_date=_date(row["Transaction Date"]), product=_text(row["Product Code"]),
            amount=_decimal(row["Net Sales"] if row["Net Sales"] is not None else _decimal(row["Gross Sales"]) - _decimal(row["Discount Amount"]) - _decimal(row["Return Amount"])),
            status=_text(row["Record Status"]).casefold(), currency=_text(row["Currency"]),
            eligibility_override=_optional_bool(row["Eligibility Override"]),
            duplicate_of=_text(row["Duplicate Of"]) or None,
            data_quality_flag=_text(row["Data Quality Flag"]), source_checksum=canonical_checksum(row),
        ) for row in transaction_rows]

        payments = [Payment(
            id=_text(row["Payment ID"]), amount=_decimal(row["Payment Amount"]),
            status=_text(row["Payment Status"]).casefold(), currency=_text(row["Currency"]),
            approval_status=_text(row["Approval Status"]).casefold() or None,
            applied_rate=_decimal(row["Applied Rate"]) if row["Applied Rate"] is not None else None,
            agreement_id_used=_text(row["Agreement ID Used"]) or None,
            adjustment_id=_text(row["Adjustment ID"]) or None,
            source_checksum=canonical_checksum(row),
        ) for row in payment_rows]

        return EvidenceBundle(
            agreements=agreements, transactions=transactions, payments=payments,
            source_systems={"agreements": "ContractHub", "transactions": "ERP-Ledger", "payments": "PaymentHub"},
            retrieved_at=datetime.now(timezone.utc),
        )
